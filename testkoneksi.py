import os
import time
from pathlib import Path

import oracledb
import openpyxl
import pandas as pd
import schedule

# ================================================================
# ORACLE CLIENT 19c / THICK MODE
# ================================================================
CLIENT_LIB_DIR = r"C:\oraclexe\instantclient_19_32"
OCI_DLL = os.path.join(CLIENT_LIB_DIR, "oci.dll")
if not os.path.isfile(OCI_DLL):
    raise FileNotFoundError(f"oci.dll tidak ditemukan: {OCI_DLL}")

oracledb.init_oracle_client(lib_dir=CLIENT_LIB_DIR)

# ================================================================
# DATABASE
# ================================================================
DB_USER = "MONPD"
DB_PASS = os.getenv("MONPD_DB_PASS", "monpd2025")
DB_HOST = "10.21.39.80"
DB_PORT = "1521"
DB_SERVICE = "DEVDB"
DB_DSN = f"{DB_HOST}:{DB_PORT}/{DB_SERVICE}"

# ================================================================
# FILE
# ================================================================
FOLDER_PATH = Path(r"D:\DATA IWAN\python stream jasnita\files")
FILE_CSV = FOLDER_PATH / "log_kendaraan_terhitung.csv"
FILE_XLSX = FOLDER_PATH / "log_capture_kendaraan.xlsx"

# ================================================================
# MAPPING TEXT CSV -> NUMBER ORACLE
# Sesuaikan ANGKA di sini dengan CHECK CONSTRAINT database Anda.
# Berdasarkan format generator CSV saat ini: MOBIL/MOTOR, IN, BAPENDA.
# ================================================================
JENIS_KEND_MAP = {
    "MOBIL": 2,
    "MOTOR": 1,
}
DIRECTION_MAP = {
    "IN": 1,
    "OUT": 2,
}
VENDOR_MAP = {
    "BAPENDA": 3,
}

# ================================================================
# UTILITAS
# ================================================================
def clean_value(value):
    """
    Membersihkan nilai tanpa mengubah representasi angka menjadi float.

    PENTING: jangan melakukan float()/int() di sini karena NOP adalah
    identifier panjang dan konversi ke float dapat kehilangan digit.
    Contoh: 357811000390700825 -> 357811000390700800.
    """
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return str(value).strip()


def clean_id(value):
    """
    Membersihkan kolom ID saja. ID boleh dinormalisasi dari 64.0 -> 64.
    Tidak digunakan untuk NOP.
    """
    text = clean_value(value)
    if not text:
        return ""
    try:
        number = float(text)
        if number.is_integer():
            return str(int(number))
    except (ValueError, TypeError):
        pass
    return text


def clean_key_text(value):
    """
    Membersihkan key text seperti NOP/CCTV_ID tanpa pernah mengubahnya
    menjadi float/int. Ini menjaga seluruh digit NOP tetap utuh.
    """
    return clean_value(value)


def nullable_text(value):
    value = clean_value(value)
    if value.upper() in {"", "-", "NULL", "NONE", "N/A"}:
        return None
    return value


def parse_datetime(value, field_name):
    value = clean_value(value)
    if not value:
        return None
    dt = pd.to_datetime(value, errors="coerce")
    if pd.isna(dt):
        raise ValueError(f"{field_name} tidak valid: {value!r}")
    try:
        if dt.tzinfo is not None:
            dt = dt.tz_localize(None)
    except Exception:
        pass
    return dt.to_pydatetime()


def parse_number_or_null(value, field_name):
    value = clean_value(value)
    if value.upper() in {"", "-", "NULL", "NONE", "N/A"}:
        return None
    try:
        number = float(value)
        if number.is_integer():
            return int(number)
        raise ValueError
    except ValueError:
        raise ValueError(f"{field_name} harus NUMBER/NULL, tetapi {value!r}")


def parse_code(value, mapping, field_name):
    value = clean_value(value)
    if value.upper() in {"", "-", "NULL", "NONE", "N/A"}:
        return None
    try:
        number = float(value)
        if number.is_integer():
            return int(number)
    except (ValueError, TypeError):
        pass
    key = value.upper()
    if key not in mapping:
        raise ValueError(
            f"{field_name}={value!r} belum mempunyai mapping. "
            f"Mapping aktif: {mapping}"
        )
    return mapping[key]


def oracle_error(exc):
    code = getattr(exc, "code", None)
    return f"ORA-{int(code):05d}: {exc}" if code is not None else str(exc)

# ================================================================
# SQL
# ================================================================
SQL_PARENT_EXISTS = """
    SELECT 1
    FROM T_OP_PARKIR_CCTV_IH
    WHERE ID = :id AND NOP = :nop AND CCTV_ID = :cctv_id
"""

SQL_PARENT_INSERT = """
    INSERT INTO T_OP_PARKIR_CCTV_IH
    (
        ID, NOP, CCTV_ID, NAMA_OP, ALAMAT_OP, WILAYAH_PAJAK,
        WAKTU_MASUK, JENIS_KEND, PLAT_NO, WAKTU_KELUAR,
        DIRECTION, LOG, IMAGE_URL, VENDOR
    )
    VALUES
    (
        :id, :nop, :cctv_id, :nama_op, :alamat_op, :wilayah_pajak,
        :waktu_masuk, :jenis_kend, :plat_no, :waktu_keluar,
        :direction, :log_data, :image_url, :vendor
    )
"""

SQL_CHILD_STATUS = """
    SELECT CASE WHEN IMAGE_DATA IS NULL THEN 0 ELSE 1 END
    FROM T_OP_PARKIR_CCTV_DOK_IH
    WHERE ID = :id AND NOP = :nop AND CCTV_ID = :cctv_id
"""

SQL_BLOB_INSERT = """
    INSERT INTO T_OP_PARKIR_CCTV_DOK_IH
    (ID, NOP, CCTV_ID, IMAGE_DATA)
    VALUES
    (:id, :nop, :cctv_id, :image_data)
"""

SQL_BLOB_UPDATE = """
    UPDATE T_OP_PARKIR_CCTV_DOK_IH
    SET IMAGE_DATA = :image_data
    WHERE ID = :id AND NOP = :nop AND CCTV_ID = :cctv_id
"""

# ================================================================
# CHECK CONSTRAINT DIAGNOSTIC
# ================================================================
def show_check_constraints(cursor):
    print("\n[DB] CHECK constraint target:")
    try:
        cursor.execute("""
            SELECT constraint_name, search_condition
            FROM user_constraints
            WHERE table_name = 'T_OP_PARKIR_CCTV_IH'
              AND constraint_name IN
                  ('SYS_C0041739','SYS_C0041740','SYS_C0041741')
            ORDER BY constraint_name
        """)
        rows = cursor.fetchall()
        for name, condition in rows:
            print(f"  {name}: {condition}")
    except Exception as exc:
        print(f"  [WARN] search_condition tidak terbaca: {oracle_error(exc)}")

# ================================================================
# PARENT CHECK / CHILD STATUS
# ================================================================
def parent_exists(cursor, id_value, nop, cctv):
    cursor.execute(SQL_PARENT_EXISTS, {
        "id": id_value, "nop": nop, "cctv_id": cctv
    })
    return cursor.fetchone() is not None


def child_status(cursor, id_value, nop, cctv):
    cursor.execute(SQL_CHILD_STATUS, {
        "id": id_value, "nop": nop, "cctv_id": cctv
    })
    row = cursor.fetchone()
    return None if row is None else int(row[0])

# ================================================================
# CSV -> PARENT
# ================================================================
def process_csv(conn, cursor):
    inserted = existing = failed = 0

    print("\n" + "=" * 90)
    print("A. CSV -> T_OP_PARKIR_CCTV_IH")
    print("=" * 90)

    if not FILE_CSV.exists():
        print(f"[CSV] File tidak ditemukan: {FILE_CSV}")
        return inserted, existing, failed

    df = pd.read_csv(FILE_CSV, dtype=str, keep_default_na=False)
    print(f"[CSV] {FILE_CSV}")
    print(f"[CSV] {len(df)} baris")
    print(f"[MAP] JENIS_KEND={JENIS_KEND_MAP}")
    print(f"[MAP] DIRECTION={DIRECTION_MAP}")
    print(f"[MAP] VENDOR={VENDOR_MAP}")

    required = [
        "ID", "NOP", "CCTV_ID", "NAMA_OP", "ALAMAT_OP", "WILAYAH_PAJAK",
        "WAKTU_MASUK", "JENIS_KEND", "PLAT_NO", "WAKTU_KELUAR",
        "DIRECTION", "LOG", "IMAGE_URL", "VENDOR"
    ]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise RuntimeError("Kolom CSV hilang: " + ", ".join(missing))

    for idx, row in df.iterrows():
        line = idx + 2
        id_value = nop = cctv = ""
        try:
            id_value = clean_value(row["ID"])
            nop = clean_value(row["NOP"])
            cctv = clean_value(row["CCTV_ID"])
            if not id_value or not nop or not cctv:
                raise ValueError("ID/NOP/CCTV_ID wajib diisi")

            # Jangan insert dua kali bila key sudah ada.
            if parent_exists(cursor, id_value, nop, cctv):
                existing += 1
                print(f"[CSV {line}] sudah ada | ID={id_value}")
                continue

            data = {
                "id": id_value,
                "nop": nop,
                "cctv_id": cctv,
                "nama_op": nullable_text(row["NAMA_OP"]),
                "alamat_op": nullable_text(row["ALAMAT_OP"]),
                "wilayah_pajak": parse_number_or_null(row["WILAYAH_PAJAK"], "WILAYAH_PAJAK"),
                "waktu_masuk": parse_datetime(row["WAKTU_MASUK"], "WAKTU_MASUK"),
                "jenis_kend": parse_code(row["JENIS_KEND"], JENIS_KEND_MAP, "JENIS_KEND"),
                "plat_no": nullable_text(row["PLAT_NO"]),
                "waktu_keluar": parse_datetime(row["WAKTU_KELUAR"], "WAKTU_KELUAR"),
                "direction": parse_code(row["DIRECTION"], DIRECTION_MAP, "DIRECTION"),
                "log_data": nullable_text(row["LOG"]),
                "image_url": nullable_text(row["IMAGE_URL"]),
                "vendor": parse_code(row["VENDOR"], VENDOR_MAP, "VENDOR"),
            }

            if data["waktu_masuk"] is None:
                raise ValueError("WAKTU_MASUK tidak boleh NULL")
            if data["jenis_kend"] is None:
                raise ValueError("JENIS_KEND tidak boleh NULL")
            if data["direction"] is None:
                raise ValueError("DIRECTION tidak boleh NULL")
            if data["vendor"] is None:
                raise ValueError("VENDOR tidak boleh NULL")

            cursor.execute(SQL_PARENT_INSERT, data)
            inserted += 1
            print(
                f"[CSV {line}] INSERT OK | ID={id_value} | "
                f"JENIS={data['jenis_kend']} | "
                f"DIRECTION={data['direction']} | VENDOR={data['vendor']}"
            )

        except Exception as exc:
            failed += 1
            print("-" * 90)
            print(f"[CSV ERROR] baris={line}")
            print(f"  ID={id_value!r}")
            print(f"  NOP={nop!r}")
            print(f"  CCTV_ID={cctv!r}")
            print(f"  ERROR={oracle_error(exc)}")
            print("-" * 90)

    conn.commit()
    print(f"[CSV] COMMIT selesai | insert={inserted}, existing={existing}, gagal={failed}")
    return inserted, existing, failed

# ================================================================
# XLSX -> BLOB
# ================================================================
def process_xlsx(conn, cursor):
    inserted = updated = existing = parent_missing = invalid = failed = 0

    print("\n" + "=" * 90)
    print("B. XLSX -> T_OP_PARKIR_CCTV_DOK_IH")
    print("=" * 90)

    if not FILE_XLSX.exists():
        print(f"[XLSX] File tidak ditemukan: {FILE_XLSX}")
        return inserted, updated, existing, parent_missing, invalid, failed

    wb = openpyxl.load_workbook(FILE_XLSX, data_only=True)
    try:
        sheet = wb.active
        images = getattr(sheet, "_images", [])
        print(f"[XLSX] ditemukan {len(images)} gambar")

        for image_no, img in enumerate(images, start=1):
            row_idx = None
            try:
                row_idx = img.anchor._from.row + 1
                id_value = clean_value(sheet.cell(row=row_idx, column=1).value)
                nop = clean_value(sheet.cell(row=row_idx, column=2).value)
                cctv = clean_value(sheet.cell(row=row_idx, column=3).value)

                print(f"\n[IMAGE #{image_no}] Excel row={row_idx}")
                print(f"  ID={id_value}")
                print(f"  NOP={nop}")
                print(f"  CCTV_ID={cctv}")

                if not id_value or not nop or not cctv:
                    print("  SKIP: key kosong")
                    invalid += 1
                    continue

                image_bytes = img._data()
                if not image_bytes:
                    print("  SKIP: image kosong")
                    invalid += 1
                    continue

                print(f"  Ukuran={len(image_bytes):,} bytes")

                # FK R01 mensyaratkan parent ID+NOP+CCTV_ID sudah ada.
                if not parent_exists(cursor, id_value, nop, cctv):
                    print("  SKIP: PARENT TIDAK DITEMUKAN")
                    parent_missing += 1
                    continue

                status = child_status(cursor, id_value, nop, cctv)

                if status is None:
                    cursor.execute(SQL_BLOB_INSERT, {
                        "id": id_value,
                        "nop": nop,
                        "cctv_id": cctv,
                        "image_data": image_bytes,
                    })
                    inserted += 1
                    print("  RESULT: BLOB INSERT BERHASIL")

                elif status == 0:
                    cursor.execute(SQL_BLOB_UPDATE, {
                        "id": id_value,
                        "nop": nop,
                        "cctv_id": cctv,
                        "image_data": image_bytes,
                    })
                    updated += 1
                    print("  RESULT: BLOB UPDATE BERHASIL")

                else:
                    existing += 1
                    print("  RESULT: IMAGE SUDAH ADA - SKIP")

            except Exception as exc:
                failed += 1
                print(f"  ERROR: {oracle_error(exc)}")

        conn.commit()
        print(
            f"[BLOB] COMMIT selesai | insert={inserted}, update={updated}, "
            f"existing={existing}, parent_missing={parent_missing}, "
            f"invalid={invalid}, gagal={failed}"
        )

    finally:
        wb.close()

    return inserted, updated, existing, parent_missing, invalid, failed

# ================================================================
# MAIN
# ================================================================
def process_and_insert():
    print("\n" + "#" * 90)
    print(f"[{time.strftime('%Y-%m-%d %H:%M:%S')}] SINKRONISASI DIMULAI")
    print("#" * 90)

    conn = cursor = None
    try:
        conn = oracledb.connect(
            user=DB_USER,
            password=DB_PASS,
            dsn=DB_DSN,
        )
        cursor = conn.cursor()

        print(f"[ORACLE] Client={oracledb.clientversion()}")
        print(f"[ORACLE] Database={conn.version}")
        print(f"[ORACLE] Thin={conn.thin}")

        if conn.thin:
            raise RuntimeError("Oracle 11.x harus menggunakan Thick Mode")

        show_check_constraints(cursor)

        parent_result = process_csv(conn, cursor)

        # Hanya proses BLOB setelah tahap parent selesai.
        blob_result = process_xlsx(conn, cursor)

        print("\n" + "=" * 90)
        print("SUMMARY")
        print("=" * 90)
        print(f"Parent INSERT       : {parent_result[0]}")
        print(f"Parent sudah ada   : {parent_result[1]}")
        print(f"Parent gagal       : {parent_result[2]}")
        print(f"BLOB INSERT        : {blob_result[0]}")
        print(f"BLOB UPDATE        : {blob_result[1]}")
        print(f"BLOB sudah ada     : {blob_result[2]}")
        print(f"BLOB parent missing: {blob_result[3]}")
        print(f"BLOB invalid       : {blob_result[4]}")
        print(f"BLOB error         : {blob_result[5]}")
        print("=" * 90)

    except Exception as exc:
        print(f"\nERROR UTAMA: {oracle_error(exc)}")
        if conn:
            try:
                conn.rollback()
                print("Transaction di-ROLLBACK.")
            except Exception:
                pass
    finally:
        if cursor:
            try:
                cursor.close()
            except Exception:
                pass
        if conn:
            try:
                conn.close()
            except Exception:
                pass
        print("Koneksi Oracle ditutup.")

# ================================================================
# SCHEDULER
# ================================================================
if __name__ == "__main__":
    process_and_insert()
    schedule.every(1).minutes.do(process_and_insert)
    print("\n[INFO] Scheduler aktif. Interval= menit. Ctrl+C untuk berhenti.")
    while True:
        try:
            schedule.run_pending()
            time.sleep(1)
        except KeyboardInterrupt:
            print("Program dihentikan.")
            break
